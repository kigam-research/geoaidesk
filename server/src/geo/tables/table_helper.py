"""Read-only bounded table overlay. Never evaluates formulas or rewrites source data."""
import csv
import io
import json
import math
import os
import re
import sys
import zipfile
from collections import defaultdict

LIMITS = dict(sourceBytes=16 * 1024 * 1024, rows=20000, columns=64,
              cells=200000, parsedBytes=4 * 1024 * 1024, issues=200)
ALIASES = {
    "hole_id": ["holeid", "hole", "borehole", "drillhole", "bhid", "wellid"],
    "sample_id": ["sampleid", "sample", "sampid"],
    "from_depth": ["from", "fromdepth", "depthfrom", "startdepth"],
    "to_depth": ["to", "todepth", "depthto", "enddepth"],
    "depth_unit": ["depthunit", "depthunits"],
    "x": ["x", "easting", "longitude", "lon"],
    "y": ["y", "northing", "latitude", "lat"],
    "z": ["z", "elevation", "rl"], "crs": ["crs", "epsg"],
    "survey_depth": ["surveydepth", "atdepth", "depth"],
    "azimuth": ["azimuth", "azim", "azi"], "dip": ["dip", "inclination"],
    "analyte": ["analyte", "element", "parameter"],
    "value": ["value", "result", "assay", "concentration"],
    "unit": ["unit", "units", "assayunit"],
    "detection_limit": ["detectionlimit", "dl", "lod", "mdl"],
    "qualifier": ["qualifier", "qual", "censor", "flag"],
}


def text(v):
    return None if v is None else str(v)


def load(src):
    filename = src["abs"]
    if os.path.getsize(filename) > LIMITS["sourceBytes"]:
        raise ValueError("Source exceeds 16 MiB")
    reasons, sheets, sheet = [], [], None
    book = None
    if filename.lower().endswith(".xlsx"):
        # Bound the decompressed container before openpyxl loads shared strings/styles.
        with zipfile.ZipFile(filename) as z:
            if len(z.infolist()) > 4096 or sum(i.file_size for i in z.infolist()) > 64 * 1024 * 1024:
                raise ValueError("XLSX expanded-size limit exceeded")
        from openpyxl import load_workbook
        book = load_workbook(filename, read_only=True, data_only=False, keep_links=False)
        if len(book.sheetnames) > 50:
            book.close()
            raise ValueError("XLSX sheet limit exceeded (50)")
        sheets = book.sheetnames
        sheet = src.get("sheet") or sheets[0]
        if sheet not in sheets:
            book.close()
            raise ValueError("Worksheet not found")
        ws = book[sheet]
        ws.reset_dimensions()  # Don't trust stale producer-provided dimension metadata.
        iterator = enumerate(ws.iter_rows(values_only=True), 1)
        stream = None
    else:
        if src.get("sheet"):
            raise ValueError("Sheets apply only to XLSX")
        # UTF-8/UTF-8-BOM only: no silent encoding substitution.
        with open(filename, "rb") as f:
            raw = f.read(LIMITS["sourceBytes"] + 1)
        if len(raw) > LIMITS["sourceBytes"]:
            raise ValueError("Source exceeds 16 MiB")
        stream = io.StringIO(raw.decode("utf-8-sig"), newline="")
        csv.field_size_limit(65536)
        reader = csv.reader(stream, delimiter="\t" if filename.lower().endswith(".tsv") else ",", strict=True)
        def records():
            while True:
                line = reader.line_num + 1
                try:
                    row = next(reader)
                except StopIteration:
                    return
                yield line, row
        iterator = records()
    try:
        first = next(iterator, None)
        if first is None:
            raise ValueError("Table has no header")
        headers = [text(v) or "" for v in first[1]]
        if not headers or not any(headers) or len(headers) > LIMITS["columns"] or any(len(h) > 256 for h in headers):
            raise ValueError("Header limit exceeded or empty header (64 columns, 256 characters each)")
        rows, size, cells = [], 0, 0
        for source_row, values in iterator:
            values = [text(v) for v in values]
            # openpyxl may pad to worksheet width. Ignore only empty trailing cells.
            while len(values) > len(headers) and values[-1] in (None, ""):
                values.pop()
            if len(values) > len(headers):
                reasons.append("row_wider_than_header")
                break
            values += [None] * (len(headers) - len(values))
            row_bytes = len(json.dumps(values, ensure_ascii=False).encode("utf-8"))
            if len(rows) >= LIMITS["rows"] or cells + len(values) > LIMITS["cells"] or size + row_bytes > LIMITS["parsedBytes"] or any(v is not None and len(v) > 4096 for v in values):
                reasons.append("data_limit")
                break
            rows.append((source_row, values))
            size += row_bytes
            cells += len(values)
        source = dict(path=src["path"])
        if sheet:
            source["sheet"] = sheet
        return dict(source=source, sheets=sheets, columns=headers, rows=rows, reasons=reasons,
                    status="partial" if reasons else "complete")
    finally:
        if book:
            book.close()
        if stream:
            stream.close()


def suggest(headers):
    candidates = {}
    for role, aliases in ALIASES.items():
        matches = [i for i, h in enumerate(headers) if re.sub(r"[^a-z0-9]", "", h.lower()) in aliases]
        if len(matches) == 1:
            candidates[role] = matches[0]
    relevant = (("hole_id" in candidates and
                 (("from_depth" in candidates and "to_depth" in candidates) or
                  ("x" in candidates and "y" in candidates) or
                  ("survey_depth" in candidates and ("azimuth" in candidates or "dip" in candidates)))) or
                ("value" in candidates and ("sample_id" in candidates or "hole_id" in candidates)))
    return bool(relevant), candidates if relevant else {}


def assess(req):
    t = load(req)
    mapping = req.get("mapping", {})
    declarations = req.get("declarations", {})
    for col in mapping.values():
        if col >= len(t["columns"]):
            raise ValueError("Mapped column is outside the source header")
    relevant, suggestions = suggest(t["columns"])
    out = dict(schemaVersion="geo-tables/v1", projectId="", source=t["source"], sheets=t["sheets"], columns=t["columns"],
               rowsAssessed=len(t["rows"]), relevant=relevant, suggestions=suggestions, mapping=mapping,
               status=t["status"], reasons=t["reasons"], checks=[], issues=[], issueCount=0,
               issuesTruncated=False, preview=[], previewTruncated=False, limits=LIMITS)
    def ref(table, i):
        return dict(**table["source"], dataRow=i + 1, sourceRow=table["rows"][i][0])
    def get(i, role):
        return t["rows"][i][1][mapping[role]] if role in mapping else None
    def declared(i, role):
        v = get(i, role)
        return v if v is not None and v.strip() else declarations.get(role)
    def norm(v):
        return v.strip().casefold() if v else ""
    def issue(code, refs, role=None, values=None):
        out["issueCount"] += 1
        if len(out["issues"]) < LIMITS["issues"]:
            row = dict(code=code, refs=refs)
            if role:
                row["role"] = role
            if values is not None:
                row["values"] = [(v or "")[:256] for v in values]
            out["issues"].append(row)
        else:
            out["issuesTruncated"] = True
    def check(name, reason=None, partial=False):
        out["checks"].append(dict(name=name, status="unassessed" if reason else "partial" if partial or t["status"] == "partial" else "complete", **({"reason": reason} if reason else {})))
    for i, (_, values) in enumerate(t["rows"][:12]):
        clipped = [v[:256] if v is not None else None for v in values]
        out["previewTruncated"] |= clipped != values
        out["preview"].append(dict(ref=ref(t, i), values=clipped))
    out["previewTruncated"] |= len(t["rows"]) > 12

    # Repeated hole IDs are expected in intervals/surveys/assays; do not call them duplicates.
    id_role = "sample_id" if "sample_id" in mapping else "hole_id" if "hole_id" in mapping and not any(r in mapping for r in ("from_depth", "to_depth", "survey_depth", "value")) else None
    if id_role:
        seen = {}
        for i in range(len(t["rows"])):
            v = get(i, id_role)
            if v is None or not v.strip():
                issue("missing_id", [ref(t, i)], id_role)
                continue
            # Long-form assays may legitimately repeat a sample across different analytes.
            key = (v, declared(i, "analyte") or "")
            if key in seen:
                issue("duplicate_id", [ref(t, seen[key]), ref(t, i)], id_role)
            else:
                seen[key] = i
        check("ids")
    else:
        check("ids", "id_mapping_required")

    conflict_holes = set()
    unit_roles = [r for r in ("depth_unit", "unit") if r in mapping or declarations.get(r)]
    missing_units = False
    for role in unit_roles:
        baseline = {}
        for i in range(len(t["rows"])):
            group = get(i, "hole_id") if role == "depth_unit" else declared(i, "analyte")
            # Without analyte, the one explicitly mapped value column is a single assay domain.
            group = group or ""
            value = declared(i, role)
            if not norm(value):
                missing_units = True
                continue
            expected = declarations.get(role)
            if expected and norm(value) != norm(expected):
                issue("unit_inconsistency", [ref(t, i)], role, [expected, value])
                if role == "depth_unit":
                    conflict_holes.add(group)
            if group in baseline and norm(value) != norm(baseline[group][1]):
                j, prior = baseline[group]
                issue("unit_inconsistency", [ref(t, j), ref(t, i)], role, [prior, value])
                if role == "depth_unit":
                    conflict_holes.add(group)
            else:
                baseline.setdefault(group, (i, value))
    check("units", "unit_mapping_required" if not unit_roles else None, partial=missing_units)
    if unit_roles and missing_units:
        out["checks"][-1]["reason"] = "missing_units"

    if all(r in mapping for r in ("hole_id", "from_depth", "to_depth")):
        groups = defaultdict(list)
        skipped = False
        uncertain_holes = set()
        for i in range(len(t["rows"])):
            hole = get(i, "hole_id")
            if not hole or not hole.strip():
                issue("missing_id", [ref(t, i)], "hole_id")
                skipped = True
                continue
            try:
                a, b = float(get(i, "from_depth")), float(get(i, "to_depth"))
                if not math.isfinite(a) or not math.isfinite(b) or a >= b:
                    raise ValueError()
            except (ValueError, TypeError):
                issue("invalid_interval", [ref(t, i)])
                uncertain_holes.add(hole)
                skipped = True
                continue
            # No cross-unit comparison or assumed unit. Same declared unit required per hole.
            unit = declared(i, "depth_unit")
            if not norm(unit) or hole in conflict_holes:
                uncertain_holes.add(hole)
                skipped = True
                continue
            groups[hole].append((a, b, i))
        for hole, intervals in groups.items():
            intervals.sort(key=lambda v: (v[0], v[1], v[2]))
            furthest = None
            for a, b, i in intervals:
                if furthest:
                    end, j = furthest
                    if a < end:
                        issue("overlap", [ref(t, j), ref(t, i)])
                    elif a > end and t["status"] == "complete" and hole not in uncertain_holes:
                        issue("gap", [ref(t, j), ref(t, i)])
                if furthest is None or b > furthest[0]:
                    furthest = (b, i)
        check("intervals", partial=skipped)
        if skipped:
            out["checks"][-1]["reason"] = "interval_rows_skipped"
    else:
        check("intervals", "interval_mapping_required")

    if req.get("join"):
        jreq = req["join"]
        if jreq["left"] not in mapping:
            raise ValueError("Map the left join role first")
        right = load(jreq)
        col = jreq["rightColumn"]
        if col >= len(right["columns"]):
            raise ValueError("Join column is outside the source header")
        out["join"] = dict(source=right["source"], columns=right["columns"], rowsAssessed=len(right["rows"]), status=right["status"])
        left_index, right_index = defaultdict(list), defaultdict(list)
        for i in range(len(t["rows"])):
            left_index[get(i, jreq["left"])].append(i)
        for i, (_, row) in enumerate(right["rows"]):
            right_index[row[col]].append(i)
        for index, table, other, code, other_table in ((left_index, t, right_index, "unmatched_left", right), (right_index, right, left_index, "unmatched_right", t)):
            # Absence cannot be established in a partially read other table.
            if other_table["status"] != "complete":
                continue
            for key, indices in index.items():
                if key is None or not key.strip() or key not in other:
                    for i in indices:
                        issue(code, [ref(table, i)], jreq["left"])
        # Duplicate reference keys make a join ambiguous, even if every key matched.
        for key, indices in right_index.items():
            if key is not None and key.strip():
                for i in indices[1:]:
                    issue("duplicate_id", [ref(right, indices[0]), ref(right, i)])
        check("join", partial=right["status"] == "partial")
        if right["status"] == "partial":
            out["reasons"].append("join_data_limit")
    else:
        check("join", "join_not_requested")
    assessed = [c for c in out["checks"] if c["status"] != "unassessed"]
    if not assessed:
        out["status"] = "unassessed"
    elif out["status"] == "partial" or any(c["status"] == "partial" for c in assessed):
        out["status"] = "partial"
    return out


if __name__ == "__main__":
    try:
        result = assess(json.loads(sys.argv[1]))
        print(json.dumps(result, ensure_ascii=False, allow_nan=False, separators=(",", ":")))
    except ImportError:
        print("Table helper dependencies are missing", file=sys.stderr)
        sys.exit(3)
    except Exception as exc:
        print(str(exc)[:400], file=sys.stderr)
        sys.exit(5)
